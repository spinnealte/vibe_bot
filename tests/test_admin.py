"""Админка без БД и Telegram: доступ, проверка полей справочников, экраны."""

import asyncio
from io import BytesIO

import pytest
from aiogram.types import User
from docx import Document

from stubbot.config import Settings
from stubbot.services.admin_courses import COURSE_SPEC
from stubbot.services.admin_dictionaries import SPECS, DictField, DictKind, FieldKind, parse_field
from stubbot.tg import admin_views, keyboards, texts
from stubbot.tg.filters import IsAdmin
from stubbot.utils.documents import file_to_html
from stubbot.utils.slug import slugify


def _settings(**overrides: object) -> Settings:
    values = {"db_user": "u", "db_password": "p", "db_name": "n", "owner_telegram_id": 1, **overrides}
    return Settings(_env_file=None, **values)


def test_admin_ids_from_env_string() -> None:
    settings = _settings(admin_ids="22, 33,")
    assert settings.admin_ids == [22, 33]
    assert settings.is_admin(1)  # owner — всегда админ
    assert settings.is_admin(33)
    assert not settings.is_admin(44)
    assert _settings(admin_ids="").admin_ids == []


def test_admin_button_only_for_admins(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(keyboards, "get_settings", lambda: _settings(admin_ids="22"))
    buttons = lambda chat_id: [b.text for row in keyboards.main_menu(chat_id).keyboard for b in row]  # noqa: E731
    assert texts.BTN_ADMIN in buttons(22)
    assert texts.BTN_ADMIN in buttons(1)
    assert texts.BTN_ADMIN not in buttons(99)


def test_is_admin_filter() -> None:
    settings = _settings(admin_ids="22")
    check = lambda user_id: asyncio.run(IsAdmin()(  # noqa: E731
        None, settings=settings, event_from_user=User(id=user_id, is_bot=False, first_name="T")))
    assert check(22) and check(1)
    assert not check(99)
    assert not asyncio.run(IsAdmin()(None, settings=settings, event_from_user=None))


@pytest.mark.parametrize(
    ("field", "raw", "expected"),
    [
        (DictField("x", FieldKind.LINE, True, 20), "  Учебный   класс ", "Учебный класс"),
        (DictField("x", FieldKind.LINE, True, 5), "слишком длинно", None),
        (DictField("x", FieldKind.LINE, True, 20), "   ", None),
        (DictField("x", FieldKind.URL, False, 100), " https://yandex.ru/maps/-/C ", "https://yandex.ru/maps/-/C"),
        (DictField("x", FieldKind.URL, False, 100), "yandex.ru/maps", None),
        (DictField("x", FieldKind.URL, False, 100), "https://a.ru/b c", None),
        (DictField("x", FieldKind.FULL_NAME, True, 300), "Иванов  Сергей", "Иванов Сергей"),
        (DictField("x", FieldKind.FULL_NAME, True, 300), "Иванов", None),
        (DictField("x", FieldKind.HTML, False, 30), "<b>Хирург</b>\nстаж 20 лет", "<b>Хирург</b>\nстаж 20 лет"),
        (DictField("x", FieldKind.NUMBER, False, min_value=1, max_value=500), " 16 ", 16),
        (DictField("x", FieldKind.NUMBER, False, min_value=1, max_value=500), "0", None),
        (DictField("x", FieldKind.NUMBER, False, min_value=1, max_value=500), "16 часов", None),
        (DictField("x", FieldKind.CHOICE, False, choices=("a",)), "a", None),  # выбор — только кнопками
    ],
)
def test_parse_field(field: DictField, raw: str, expected: str | None) -> None:
    assert parse_field(field, raw) == expected


def _buttons(markup) -> list[str]:
    return [b.text for row in markup.inline_keyboard for b in row]


def test_dict_list_pages_and_archive_mark() -> None:
    spec = SPECS[DictKind.LECTURERS]  # длинные названия — по одной кнопке в ряд, 8 на странице
    items = [(i, f"Лектор {i}", True) for i in range(1, 10)] + [(10, "Старый", False)]
    text, markup = admin_views.dict_list(spec, items, page=0)
    assert "Всего: <b>10</b>" in text and "в архиве: 1" in text
    first = _buttons(markup)
    assert first.count("➕ Добавить") == 1 and "1/2" in first and "▶️" in first and "◀️" not in first
    assert [len(row) for row in markup.inline_keyboard[:8]] == [1] * 8
    _, markup = admin_views.dict_list(spec, items, page=1)
    second = _buttons(markup)
    assert "🗄 Старый" in second and "◀️" in second and "▶️" not in second
    text, markup = admin_views.dict_list(spec, [], page=0)
    assert "Пока пусто" in text and "1/1" not in _buttons(markup)


@pytest.mark.parametrize("kind", [DictKind.SPECIALTIES, DictKind.POSITIONS])
def test_short_dictionaries_listed_in_two_columns(kind: DictKind) -> None:
    """Специальности (направления) и должности — кнопками в два столбца, как в пет-проекте: 16 на странице."""
    items = [(i, f"Запись {i}", True) for i in range(1, 14)]
    _, markup = admin_views.dict_list(SPECS[kind], items, page=0)
    assert [len(row) for row in markup.inline_keyboard[:7]] == [2, 2, 2, 2, 2, 2, 1]  # 13 записей — одна страница
    assert [b.text for b in markup.inline_keyboard[0]] == ["Запись 1", "Запись 2"]
    assert "1/1" not in _buttons(markup) and "▶️" not in _buttons(markup)
    _, markup = admin_views.dict_list(SPECS[kind], items + [(i, f"Запись {i}", True) for i in range(14, 20)], page=1)
    assert [b.text for b in markup.inline_keyboard[0]] == ["Запись 17", "Запись 18"] and "2/2" in _buttons(markup)


def test_card_escapes_values_and_keeps_bio_html() -> None:
    spec = SPECS[DictKind.LECTURERS]
    values = {"full_name": "Иванов <Сергей>", "regalia": None, "bio_html": "<b>Хирург</b>"}
    text, markup = admin_views.card(spec, 5, values, active=False)
    assert "Иванов &lt;Сергей&gt;" in text
    assert "<b>Регалии:</b> —" in text
    assert "<b>Биография:</b> \n<b>Хирург</b>" in text  # HTML из Telegram — как есть
    assert texts.ADMIN_CARD_ARCHIVED in text
    buttons = _buttons(markup)
    assert "✏️ ФИО" in buttons and texts.BTN_ADMIN_RESTORE in buttons and texts.BTN_ADMIN_ARCHIVE not in buttons


def test_field_prompt_buttons() -> None:
    spec = SPECS[DictKind.VENUES]
    required, optional = spec.field("name"), spec.field("map_url")
    assert _buttons(admin_views.field_prompt(spec, required, "h", None)[1]) == [texts.BTN_ADMIN_CANCEL]
    assert _buttons(admin_views.field_prompt(spec, optional, "h", None)[1]) == [texts.BTN_ADMIN_SKIP,
                                                                                 texts.BTN_ADMIN_CANCEL]
    # Правка: «Очистить» — только если сейчас что-то есть.
    assert _buttons(admin_views.field_prompt(spec, optional, "h", "https://a.ru", item_id=3)[1]) == [
        texts.BTN_ADMIN_CLEAR, texts.BTN_ADMIN_CANCEL]
    assert _buttons(admin_views.field_prompt(spec, optional, "h", None, item_id=3)[1]) == [texts.BTN_ADMIN_CANCEL]


def test_course_choice_and_multi_prompts() -> None:
    spec = COURSE_SPEC
    text, markup = admin_views.field_prompt(spec, spec.field("default_format"), "h", "offline", item_id=7)
    assert "✅ 🏫 Очно" in _buttons(markup) and "💻 Онлайн" in _buttons(markup)
    assert texts.BTN_ADMIN_CLEAR not in _buttons(markup)  # формат обязателен — очищать нельзя
    _, markup = admin_views.field_prompt(spec, spec.field("level"), "h", None)
    assert texts.BTN_ADMIN_SKIP in _buttons(markup)  # уровень необязателен
    text, markup = admin_views.field_prompt(spec, spec.field("lecturers"), "h", None,
                                            options=[(1, "Иванов С. П."), (2, "Смирнова А. О.")], selected=[2])
    assert _buttons(markup) == ["Иванов С. П.", "✅ Смирнова А. О.", texts.BTN_ADMIN_DONE, texts.BTN_ADMIN_CANCEL]
    assert "Лекторы" in text and "необязательно" not in text
    # У курса своя подсказка к «Названию» — не та, что у специальностей.
    title_hint = admin_views.field_prompt(spec, spec.field("title"), "h", None)[0]
    assert "в афише" in title_hint and "при регистрации" not in title_hint


@pytest.mark.parametrize(
    ("title", "expected"),
    [
        ("Имплантация: от планирования до протезирования", "implantatsiya-ot-planirovaniya-do-protezirovaniya"),
        ("Эндодонтия под микроскопом!", "endodontiya-pod-mikroskopom"),
        ("FP1 Concept — Digital", "fp1-concept-digital"),
        ("«»!!!", "course"),
    ],
)
def test_slugify(title: str, expected: str) -> None:
    assert slugify(title) == expected
    assert len(slugify("очень длинное название " * 10)) <= 50


def test_docx_and_txt_to_html() -> None:
    doc = Document()
    doc.add_heading("Модуль 1. Хирургия", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("Важно: ").bold = True
    paragraph.add_run("работа <на фантомах>").italic = True
    doc.add_paragraph("Разбор случаев", style="List Bullet")
    buffer = BytesIO()
    doc.save(buffer)
    html = file_to_html("Программа.DOCX", buffer.getvalue())
    assert html == ("<b>Модуль 1. Хирургия</b>\n<b>Важно: </b><i>работа &lt;на фантомах&gt;</i>\n• Разбор случаев")
    assert file_to_html("p.txt", "День 1\r\n<теория>".encode("cp1251")) == "День 1\n&lt;теория&gt;"
    assert file_to_html("p.pdf", b"%PDF") is None
    assert file_to_html("broken.docx", b"not a zip") is None


def test_carousel_edit_button_only_for_admin() -> None:
    def carousel(admin: bool) -> list[str]:
        return _buttons(keyboards.course_carousel(0, 3, program_id=5, session_id=None, apply_text=None, notify=False,
                                                  has_program_text=True, admin=admin))

    assert texts.BTN_COURSE_EDIT in carousel(True)
    assert texts.BTN_COURSE_EDIT not in carousel(False)
